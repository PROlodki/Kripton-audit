"""
Экспорт отчёта в файл (Excel, PDF, Word).
"""
from io import BytesIO


def export_report(report, fmt):
    """
    Возвращает (content_type, filename, content: bytes).
    fmt: 'xlsx' | 'pdf' | 'docx'
    """
    fmt = (fmt or 'xlsx').lower()
    if fmt == 'xlsx':
        return _export_xlsx(report)
    if fmt == 'pdf':
        return _export_pdf(report)
    if fmt == 'docx':
        return _export_docx(report)
    return _export_xlsx(report)


def _export_xlsx(report):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment
    except ImportError:
        return _export_csv_fallback(report)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Отчет'
    ws['A1'] = report.title
    ws['A1'].font = Font(bold=True)
    row = 2
    ws.cell(row=row, column=1, value='Тип отчета:')
    ws.cell(row=row, column=2, value=report.report_type.name if report.report_type_id else '')
    row += 1
    ws.cell(row=row, column=1, value='Описание:')
    ws.cell(row=row, column=2, value=report.description or '')
    row += 2
    if report.data and isinstance(report.data, dict):
        for k, v in report.data.items():
            ws.cell(row=row, column=1, value=str(k))
            ws.cell(row=row, column=2, value=str(v) if not isinstance(v, (dict, list)) else str(v))
            row += 1
    elif report.data and isinstance(report.data, list):
        for i, item in enumerate(report.data[:500]):
            if isinstance(item, dict):
                for k, v in item.items():
                    ws.cell(row=row, column=1, value=str(k))
                    ws.cell(row=row, column=2, value=str(v)[:32767])
                    row += 1
                row += 1
            else:
                ws.cell(row=row, column=1, value=str(item)[:32767])
                row += 1
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"report_{report.id}_{report.title[:30].replace(' ', '_')}.xlsx"
    return (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        name,
        buf.getvalue(),
    )


def _export_csv_fallback(report):
    import csv
    buf = BytesIO()
    writer = csv.writer(buf)
    writer.writerow([report.title])
    writer.writerow(['Тип', report.report_type.name if report.report_type_id else ''])
    writer.writerow(['Описание', report.description or ''])
    if report.data and isinstance(report.data, list) and report.data and isinstance(report.data[0], dict):
        keys = list(report.data[0].keys())
        writer.writerow(keys)
        for item in report.data[:1000]:
            writer.writerow([item.get(k, '') for k in keys])
    buf.seek(0)
    name = f"report_{report.id}.csv"
    return 'text/csv', name, buf.getvalue()


def _export_pdf(report):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm
    except ImportError:
        content = f"Report: {report.title}\n\n{report.description or ''}\n".encode('utf-8')
        return 'text/plain', f'report_{report.id}.txt', content
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica", 14)
    c.drawString(20 * mm, 270 * mm, report.title[:80])
    c.setFont("Helvetica", 10)
    y = 250 * mm
    c.drawString(20 * mm, y, f"Type: {report.report_type.name if report.report_type_id else '-'}")
    y -= 6 * mm
    for line in (report.description or '')[:500].split('\n')[:20]:
        c.drawString(20 * mm, y, line[:90])
        y -= 5 * mm
    c.save()
    buf.seek(0)
    return 'application/pdf', f'report_{report.id}.pdf', buf.getvalue()


def _export_docx(report):
    try:
        from docx import Document
    except ImportError:
        content = f"Report: {report.title}\n\n{report.description or ''}\n".encode('utf-8')
        return 'text/plain', f'report_{report.id}.txt', content
    doc = Document()
    doc.add_heading(report.title, 0)
    doc.add_paragraph(f"Тип отчета: {report.report_type.name if report.report_type_id else '-'}")
    doc.add_paragraph(report.description or '')
    if report.data and isinstance(report.data, dict):
        for k, v in report.data.items():
            doc.add_paragraph(f"{k}: {v}")
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', f'report_{report.id}.docx', buf.getvalue()
